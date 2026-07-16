"""Unit tests for the merchant orders Excel export (ExportService).

Covers the pure row-mapping (Order → export row) and the xlsx byte builder.
The money-path correctness here is field selection + MSK date conversion, so
both are asserted explicitly.
"""
from datetime import datetime, timezone
from decimal import Decimal
from io import BytesIO

from unittest.mock import MagicMock

import openpyxl
import pytest

from app.common.enums.orders import OrderStatus
from app.modules.orders.service import ExportService


def _svc():
    # Pure row/xlsx methods never touch the session.
    return ExportService(session=None)


def _order(**over):
    o = MagicMock()
    o.merchant_id = 42
    o.external_id = "REQ-001"
    o.uuid = "62e75499-a9fd-4bcc-8c5c-aaeec8a6f611"
    o.amount = Decimal("4900.00")
    o.amount_usdt = Decimal("65.0123")
    o.status = OrderStatus.SUCCESS
    o.exchange_rate = Decimal("75.36")
    o.fee_usdt = Decimal("4.6421")
    # 2026-06-11 11:29:26 UTC  → 14:29:26 МСК
    o.created_at = datetime(2026, 6, 11, 11, 29, 26, tzinfo=timezone.utc)
    o.confirmed_at = datetime(2026, 6, 11, 11, 57, 23, tzinfo=timezone.utc)
    for k, v in over.items():
        setattr(o, k, v)
    return o


def test_row_has_one_value_per_header():
    row = _svc().order_to_export_row(_order())
    assert len(row) == len(ExportService.EXPORT_HEADERS)


def test_row_starts_with_external_id():
    row = _svc().order_to_export_row(_order())
    # Column 0 = external_id (terminal_id column was removed).
    assert row[0] == "REQ-001"


def test_row_uuid_amounts_status_rate_fee():
    row = _svc().order_to_export_row(_order())
    assert row[1] == "62e75499-a9fd-4bcc-8c5c-aaeec8a6f611"
    assert row[2] == 4900.0          # amount RUB
    assert row[3] == 65.0123          # amount USDT
    assert row[4] == "success"        # status value
    assert row[5] == 75.36            # exchange rate
    assert row[6] == 4.6421           # service commission (fee_usdt)


def test_row_created_date_converted_to_msk_naive():
    row = _svc().order_to_export_row(_order())
    # created_at 11:29 UTC → 14:29 МСК, tz stripped (openpyxl can't write aware).
    assert row[7] == datetime(2026, 6, 11, 14, 29, 26)
    assert row[7].tzinfo is None


def test_row_handles_null_optional_numbers():
    row = _svc().order_to_export_row(
        _order(amount_usdt=None, exchange_rate=None, fee_usdt=None)
    )
    assert row[3] is None
    assert row[5] is None
    assert row[6] is None


def test_build_xlsx_is_readable_with_header_and_rows():
    content = _svc().build_orders_xlsx([_order(), _order(external_id="REQ-002")])
    assert isinstance(content, bytes) and content[:2] == b"PK"  # xlsx = zip
    wb = openpyxl.load_workbook(BytesIO(content))
    ws = wb.active
    # Header row + 2 data rows.
    assert ws.max_row == 3
    header = [c.value for c in ws[1]]
    assert header == ExportService.EXPORT_HEADERS
    assert ws.cell(row=2, column=1).value == "REQ-001"   # external_id (now first col)
    assert ws.cell(row=3, column=1).value == "REQ-002"


def test_build_xlsx_empty_still_has_header():
    content = _svc().build_orders_xlsx([])
    wb = openpyxl.load_workbook(BytesIO(content))
    ws = wb.active
    assert ws.max_row == 1
    assert [c.value for c in ws[1]] == ExportService.EXPORT_HEADERS


@pytest.mark.asyncio
async def test_export_orders_xlsx_delegates_fetch_and_builds(monkeypatch):
    """The full operation fetches via OrderService.list_orders_for_export and
    serialises the result."""
    from app.modules.orders import service as svc_mod

    captured = {}

    async def fake_list(self, user_id, *, date_from, date_to, limit):
        captured.update(user_id=user_id, date_from=date_from, date_to=date_to, limit=limit)
        return [_order(), _order(external_id="REQ-002")]

    monkeypatch.setattr(svc_mod.OrderService, "list_orders_for_export", fake_list)

    df = datetime(2026, 6, 1, tzinfo=timezone.utc)
    dt = datetime(2026, 6, 30, tzinfo=timezone.utc)
    content = await ExportService(session=MagicMock()).export_orders_xlsx(
        7, date_from=df, date_to=dt, limit=100_000,
    )
    assert captured == {"user_id": 7, "date_from": df, "date_to": dt, "limit": 100_000}
    wb = openpyxl.load_workbook(BytesIO(content))
    assert wb.active.max_row == 3  # header + 2 orders
